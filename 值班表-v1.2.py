import os
import random
import re
import shutil
import threading
import queue
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import pandas as pd
import openpyxl
from openpyxl.utils import get_column_letter
from datetime import datetime

class ScheduleApp:
    def __init__(self, root):
        self.root = root
        self.root.title("校团委排班系统")
        self.root.geometry("700x500")
        
        self.file_c_path = None
        self.task_queue = queue.Queue()
        
        self.setup_ui()
        self.check_queue()

    def setup_ui(self):
        # 顶部控制区
        top_frame = tk.Frame(self.root, pady=10)
        top_frame.pack(fill=tk.X, padx=10)
        
        self.btn_select_c = tk.Button(top_frame, text="选择之前的排版文件文件(C)", command=self.select_file_c)
        self.btn_select_c.pack(side=tk.LEFT, padx=5)
        
        self.lbl_file_c = tk.Label(top_frame, text="未选择文件", fg="red")
        self.lbl_file_c.pack(side=tk.LEFT, padx=5)
        
        self.btn_start = tk.Button(top_frame, text="开始排班", command=self.start_processing, state=tk.DISABLED)
        self.btn_start.pack(side=tk.RIGHT, padx=5)
        
        # 进度条
        self.progress = ttk.Progressbar(self.root, orient="horizontal", mode="determinate")
        self.progress.pack(fill=tk.X, padx=10, pady=5)
        
        # 实时表格展示区 (对应模板 B3到F6，共4行5列)
        table_frame = tk.Frame(self.root, pady=10)
        table_frame.pack(fill=tk.BOTH, expand=True, padx=10)
        
        # 创建Treeview展示表格
        columns = ["B", "C", "D", "E", "F"]
        self.tree = ttk.Treeview(table_frame, columns=columns, show="headings", height=4)
        self.tree.pack(fill=tk.BOTH, expand=True)
        
        for col in columns:
            self.tree.heading(col, text=f"{col}列")
            self.tree.column(col, width=100, anchor="center")
            
        # 插入初始空行 (代表第3到6行)
        self.tree_rows = []
        for i in range(4):
            row = self.tree.insert("", tk.END, values=("", "", "", "", ""))
            self.tree_rows.append(row)
            
        # 底部日志区
        self.log_text = tk.Text(self.root, height=8, state=tk.DISABLED)
        self.log_text.pack(fill=tk.X, padx=10, pady=5)
        
    def log(self, message):
        """线程安全的日志输出"""
        self.task_queue.put(("log", message))
        
    def update_table(self, r_idx, c_idx, text):
        """线程安全的表格更新 (r_idx 为 0-3, c_idx 为 0-4)"""
        self.task_queue.put(("update_table", (r_idx, c_idx, text)))
        
    def update_progress(self, value):
        """线程安全的进度条更新"""
        self.task_queue.put(("progress", value))

    def select_file_c(self):
        path = filedialog.askopenfilename(title="选择C文件", filetypes=[("Excel files", "*.xlsx")])
        if path:
            self.file_c_path = path
            self.lbl_file_c.config(text=os.path.basename(path), fg="green")
            self.btn_start.config(state=tk.NORMAL)

    def start_processing(self):
        if not self.file_c_path:
            return
        self.btn_start.config(state=tk.DISABLED)
        self.btn_select_c.config(state=tk.DISABLED)
        self.progress["value"] = 0
        # 清空表格
        for row in self.tree_rows:
            self.tree.item(row, values=("", "", "", "", ""))
            
        # 开启多线程处理，防止GUI卡死
        threading.Thread(target=self.process_task, daemon=True).start()

    def process_task(self):
        try:
            self.log("开始读取数据...")
            base_file = "空课表&部门人员.xlsx"
            
            if not os.path.exists(base_file):
                self.log(f"错误: 找不到文件 {base_file}")
                return

            # 1. 读取 "部门" 工作表
            df_dept = pd.read_excel(base_file, sheet_name="部门", header=None)
            dept_dict = {}
            for index, row in df_dept.iterrows():
                try:
                    dept = str(row[0]).strip()
                    name = str(row[2]).strip()
                    if name != 'nan' and dept != 'nan':
                        dept_dict[name] = dept
                except:
                    continue
                    
            # 2. 读取 "空课表" 工作表 (范围 B2:F5)
            free_time_dict = {}
            wb_base = openpyxl.load_workbook(base_file, data_only=True)
            ws_free = wb_base["空课表"]
            
            # 读取 B2 到 F5
            for r in range(2, 6):
                for c in range(2, 7):
                    cell_coord = f"{get_column_letter(c)}{r}"
                    val = ws_free[cell_coord].value
                    if val:
                        names = re.split(r'[\s,，、\n]+', str(val).strip())
                        free_time_dict[cell_coord] = [n for n in names if n]

            # 3. 读取C文件，统计A里人物出现的次数
            self.log("正在统计目标文件中的出现次数...")
            df_c = pd.read_excel(self.file_c_path, sheet_name=0)
            flat_c = df_c.stack().value_counts().to_dict()
            
            person_counts = []
            for name, dept in dept_dict.items():
                count = flat_c.get(name, 0)
                person_counts.append({'name': name, 'dept': dept, 'count': count})
                
            # 4. 生成列表B：按次数升序排序，同次数随机
            df_counts = pd.DataFrame(person_counts)
            grouped = df_counts.groupby('count')
            list_b = []
            for count, group in grouped:
                names_in_group = group['name'].tolist()
                random.shuffle(names_in_group) 
                list_b.extend(names_in_group)

            # 5. 复制模板，生成新工作簿
            self.log("正在生成排班表...")
            current_time = datetime.now().strftime("%Y%m%d%H%M%S")
            # 修改点1：文件名使用减号
            out_file = f"排班表-{current_time}.xlsx"
            
            temp_file = f"temp_{current_time}.xlsx"
            shutil.copy(base_file, temp_file)
            wb_out = openpyxl.load_workbook(temp_file)
            
            for sheet_name in wb_out.sheetnames:
                if sheet_name != "模板":
                    del wb_out[sheet_name]
            
            ws_out = wb_out["模板"]

            # 6. 遍历模板 B3 到 F6 填表
            total_cells = 4 * 5
            current_step = 0
            
            for c in range(2, 7): # B到F列
                for r in range(3, 7): # 3到6行
                    # 修改点2 & 3：模板坐标与空课表坐标映射
                    template_coord = f"{get_column_letter(c)}{r}"
                    # 模板 r 行对应空课表 r-1 行 (例如 B3 对应 B2)
                    free_coord = f"{get_column_letter(c)}{r-1}"
                    
                    available_people = free_time_dict.get(free_coord, [])
                    
                    selected_person = None
                    for person in list_b:
                        if person in available_people:
                            selected_person = person
                            break
                    
                    if selected_person:
                        dept = dept_dict[selected_person]
                        fill_text = f"{selected_person}（校团委{dept}）"
                        ws_out[template_coord] = fill_text
                        list_b.remove(selected_person)
                        
                        # 修改点4：GUI 表格展示映射 (r=3 时对应 GUI 第0行，即 r-3)
                        self.update_table(r - 3, c - 2, fill_text)
                        self.log(f"填入 {template_coord}: {fill_text}")
                    else:
                        self.log(f"警告: {template_coord} 没有找到合适的空闲人员")
                        self.update_table(r - 3, c - 2, "无合适人选")
                    
                    current_step += 1
                    self.update_progress((current_step / total_cells) * 100)

            wb_out.save(out_file)
            wb_out.close()
            os.remove(temp_file)
            
            self.log(f"排班完成！文件已保存为: {out_file}")
            self.task_queue.put(("done", None))

        except Exception as e:
            self.log(f"发生错误: {str(e)}")
            self.task_queue.put(("error", str(e)))
            self.task_queue.put(("done", None))

    def check_queue(self):
        """主线程定期检查队列，更新GUI"""
        try:
            while True:
                task_type, data = self.task_queue.get_nowait()
                
                if task_type == "log":
                    self.log_text.config(state=tk.NORMAL)
                    self.log_text.insert(tk.END, data + "\n")
                    self.log_text.see(tk.END)
                    self.log_text.config(state=tk.DISABLED)
                elif task_type == "update_table":
                    r, c, text = data
                    row_id = self.tree_rows[r]
                    current_values = list(self.tree.item(row_id, "values"))
                    current_values[c] = text
                    self.tree.item(row_id, values=current_values)
                elif task_type == "progress":
                    self.progress["value"] = data
                elif task_type == "done":
                    self.btn_start.config(state=tk.NORMAL)
                    self.btn_select_c.config(state=tk.NORMAL)
                elif task_type == "error":
                    messagebox.showerror("错误", data)
                    
        except queue.Empty:
            pass
        
        self.root.after(100, self.check_queue)

if __name__ == "__main__":
    root = tk.Tk()
    app = ScheduleApp(root)
    root.mainloop()
